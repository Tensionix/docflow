#!/usr/bin/env python3
"""Build Word templates, samples and the review description from formatting standards.

A standard is a YAML file in config/standards: the page, the corporate style set and
the numbering, each value with its basis - the house standard, a GOST clause or a
proposal awaiting review. One file gives the template (.dotx) Word creates new
documents from, a sample document that shows every style at work, and the
description colleagues correct. Corrections go back into the file, never into the
built documents.
"""

from __future__ import annotations

# >>> audion CLI bootstrap >>>
import os as _os, sys as _sys
_HERE = _os.path.dirname(_os.path.abspath(__file__))
if _HERE not in _sys.path:
    _sys.path.insert(0, _HERE)
try:
    _sys.stdout.reconfigure(encoding="utf-8")
    _sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass
# <<< audion CLI bootstrap <<<

import argparse
from dataclasses import dataclass, field
from pathlib import Path
import re
import zipfile
from xml.sax.saxutils import escape

import yaml
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
W = "{%s}" % W_NS
NSMAP = {"w": W_NS, "r": R_NS}
BACKSLASH = chr(92)

TWIPS_PER_CM = 1440 / 2.54
A4_TWIPS = (11906, 16838)
A4_WIDTH_MM = 210
CORPORATE_STYLE_RANGE = (20, 40)

TEMPLATE_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
DOCUMENT_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"

# Word keeps built-in styles under English names and shows them translated; a style
# written as "Заголовок 1" would become a second, custom heading.
BUILTIN_STYLES = {
    "Обычный": ("Normal", "Normal"),
    "Заголовок оглавления": ("TOC Heading", "TOCHeading"),
    "Название объекта": ("caption", "Caption"),
    "Абзац списка": ("List Paragraph", "ListParagraph"),
    "Маркированный список": ("List Bullet", "ListBullet"),
    "Маркированный список 2": ("List Bullet 2", "ListBullet2"),
    "Нумерованный список": ("List Number", "ListNumber"),
    "Нумерованный список 2": ("List Number 2", "ListNumber2"),
    "Список литературы": ("Bibliography", "Bibliography"),
    "Верхний колонтитул": ("header", "Header"),
    "Нижний колонтитул": ("footer", "Footer"),
    "Текст сноски": ("footnote text", "FootnoteText"),
    "Знак сноски": ("footnote reference", "FootnoteReference"),
    "Гиперссылка": ("Hyperlink", "Hyperlink"),
    "Номер страницы": ("page number", "PageNumber"),
    "Сетка таблицы": ("Table Grid", "TableGrid"),
}
for _level in range(1, 10):
    BUILTIN_STYLES["Заголовок %d" % _level] = ("heading %d" % _level, "Heading%d" % _level)
    BUILTIN_STYLES["Оглавление %d" % _level] = ("toc %d" % _level, "TOC%d" % _level)

# Styles every Word document declares as defaults; they count toward the style set.
WORD_DEFAULT_STYLES = ("Основной шрифт абзаца", "Обычная таблица", "Нет списка")
# Styles that stay out of the quick style gallery: Word applies them itself.
OUTSIDE_GALLERY = {"Верхний колонтитул", "Нижний колонтитул", "Текст сноски", "Абзац списка",
                   "Заголовок оглавления", "Оглавление 1", "Оглавление 2", "Оглавление 3"}

PAGE_KEYS = {"size": "text", "margins_mm": "margins", "header_mm": "number", "footer_mm": "number",
             "page_number": "text", "first_page_number": "flag"}
DEFAULT_KEYS = {"font": "text", "size": "number", "line": "number"}
STYLE_KEYS = {
    "font": "text", "size": "number", "bold": "flag", "italic": "flag", "caps": "flag",
    "underline": "flag", "color": "text", "vert_align": "text", "align": "text",
    "first_line": "number", "left": "number", "line": "number", "before": "number",
    "after": "number", "tabs": "text", "numbering": "text", "keep_next": "flag",
    "keep_lines": "flag", "page_break_before": "flag", "no_hyphenation": "flag",
    "outline": "integer", "borders": "text",
}
STYLE_META = {"name", "type", "role", "based_on", "next", "basis", "note", "rules"}
STYLE_TYPES = ("paragraph", "character", "table")
ALIGNMENTS = {"left": "left", "center": "center", "right": "right", "both": "both"}
TAB_KINDS = {"left": "по левому краю", "center": "по центру", "right": "по правому краю"}
NUMBER_FORMATS = {"decimal": "1", "russianLower": "а", "lowerLetter": "a", "bullet": ""}

TRANSLIT = dict(zip(
    "абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
    ["a", "b", "v", "g", "d", "e", "e", "zh", "z", "i", "j", "k", "l", "m", "n", "o", "p", "r",
     "s", "t", "u", "f", "h", "c", "ch", "sh", "shh", "", "y", "", "e", "ju", "ja"],
))


class StandardError(Exception):
    """A standard that cannot be built; the message names the file, the style and the key."""


@dataclass
class Value:
    value: object
    basis: str


@dataclass
class Style:
    name: str
    type: str
    role: str
    based_on: str | None
    next: str | None
    note: str
    props: dict[str, Value]
    rules: list[tuple[str, str, str]]
    word_name: str = ""
    style_id: str = ""
    builtin: bool = False


@dataclass
class Standard:
    path: Path
    id: str
    title: str
    lead: str
    template: str
    sample: str
    standard_basis: str
    gost_name: str
    page: dict[str, Value]
    page_rules: list[tuple[str, str, str]]
    defaults: dict[str, Value]
    numbering: dict[str, list[dict]]
    styles: list[Style] = field(default_factory=list)

    @property
    def by_name(self) -> dict[str, Style]:
        return {style.name: style for style in self.styles}

    @property
    def numbering_ids(self) -> dict[str, int]:
        return {list_id: index + 1 for index, list_id in enumerate(self.numbering)}

    @property
    def style_count(self) -> int:
        return len(self.styles) + len(WORD_DEFAULT_STYLES)

    def text_width_cm(self) -> float:
        _, _, left, right = self.page["margins_mm"].value
        return (A4_WIDTH_MM - left - right) / 10


# ----------------------------------------------------------------- loading

def style_id_for(name: str) -> str:
    words = [word for word in re.split(r"[^0-9A-Za-zА-Яа-яЁё]+", name) if word]
    latin = ["".join(TRANSLIT.get(char.lower(), char) for char in word) for word in words]
    return "".join(word[:1].upper() + word[1:] for word in latin) or "Style"


def parse_basis(raw: str, standard_basis: str, where: str) -> str:
    kind, _, rest = raw.strip().partition(" ")
    if kind == "gost" and rest.strip():
        return "gost " + rest.strip()
    if kind == "proposal" and not rest.strip():
        return "proposal"
    if kind == "standard" and not rest.strip():
        if not standard_basis:
            raise StandardError("%s: основание standard, но в файле нет standard_basis" % where)
        return "standard"
    raise StandardError("%s: неизвестное основание %r (gost <пункты>, standard или proposal)" % (where, raw))


def split_value(raw: object, default_basis: str, standard_basis: str, where: str) -> tuple[object, str]:
    if isinstance(raw, str) and "@" in raw:
        value, _, basis = raw.rpartition("@")
        return value.strip(), parse_basis(basis, standard_basis, where)
    return raw, default_basis


def convert(kind: str, raw: object, where: str) -> object:
    text = str(raw).strip()
    if kind == "text":
        return text
    if kind == "flag":
        if isinstance(raw, bool):
            return raw
        if text.lower() in ("yes", "true", "да"):
            return True
        if text.lower() in ("no", "false", "нет"):
            return False
        raise StandardError("%s: нужно yes или no, а не %r" % (where, raw))
    if kind in ("number", "integer"):
        try:
            number = float(text.replace(",", "."))
        except ValueError:
            raise StandardError("%s: нужно число, а не %r" % (where, raw)) from None
        return int(number) if kind == "integer" else number
    if kind == "margins":
        parts = text.split()
        if len(parts) != 4:
            raise StandardError("%s: поля пишутся четырьмя числами: верх низ лево право" % where)
        return tuple(convert("number", part, where) for part in parts)
    raise StandardError("%s: неизвестный тип %s" % (where, kind))


def read_values(data: dict, keys: dict[str, str], default_basis: str, standard_basis: str, where: str) -> dict[str, Value]:
    values = {}
    for key, raw in data.items():
        if key in ("basis", "rules"):
            continue
        if key not in keys:
            raise StandardError("%s: неизвестный параметр %r; допустимы: %s" % (where, key, ", ".join(sorted(keys))))
        value, basis = split_value(raw, default_basis, standard_basis, "%s.%s" % (where, key))
        values[key] = Value(convert(keys[key], value, "%s.%s" % (where, key)), basis)
    return values


def read_rules(items: list | None, standard_basis: str, where: str) -> list[tuple[str, str, str]]:
    rules = []
    for index, item in enumerate(items or [], 1):
        parts = [part.strip() for part in str(item).split(" | ")]
        if len(parts) != 3 or not all(parts):
            raise StandardError("%s.rules[%d]: строка пишется так: 'параметр | значение | основание'" % (where, index))
        rules.append((parts[0], parts[1], parse_basis(parts[2], standard_basis, "%s.rules[%d]" % (where, index))))
    return rules


def load_standard(path: Path) -> Standard:
    path = Path(path)
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except yaml.YAMLError as exc:
        raise StandardError("%s: YAML не читается: %s" % (path.name, exc)) from None
    where = path.name
    for key in ("id", "title", "template", "sample", "page", "defaults", "styles"):
        if key not in data:
            raise StandardError("%s: нет раздела %r" % (where, key))
    standard_basis = str(data.get("standard_basis") or "")
    page_data = data["page"] or {}
    page_basis = parse_basis(str(page_data.get("basis", "proposal")), standard_basis, where + ".page")
    standard = Standard(
        path=path,
        id=str(data["id"]),
        title=str(data["title"]),
        lead=str(data.get("lead") or ""),
        template=str(data["template"]),
        sample=str(data["sample"]),
        standard_basis=standard_basis,
        gost_name=str(data.get("gost_name") or "ГОСТ"),
        page=read_values(page_data, PAGE_KEYS, page_basis, standard_basis, where + ".page"),
        page_rules=read_rules(page_data.get("rules"), standard_basis, where + ".page"),
        defaults=read_values(data["defaults"] or {}, DEFAULT_KEYS, "proposal", standard_basis, where + ".defaults"),
        numbering={str(key): list(levels or []) for key, levels in (data.get("numbering") or {}).items()},
    )
    for key in ("size", "margins_mm"):
        if key not in standard.page:
            raise StandardError("%s.page: нет параметра %r" % (where, key))
    if str(standard.page["size"].value).upper() != "A4":
        raise StandardError("%s.page.size: сборщик пока знает только лист A4" % where)
    for list_id, levels in standard.numbering.items():
        if not levels:
            raise StandardError("%s.numbering.%s: нет уровней" % (where, list_id))
        for number, level in enumerate(levels, 1):
            if level.get("format") not in NUMBER_FORMATS or not str(level.get("text", "")):
                raise StandardError("%s.numbering.%s[%d]: нужны format (%s) и text" % (where, list_id, number, ", ".join(NUMBER_FORMATS)))

    seen_ids: dict[str, str] = {}
    for index, item in enumerate(data["styles"] or [], 1):
        if not isinstance(item, dict) or not item.get("name"):
            raise StandardError("%s.styles[%d]: у стиля нет name" % (where, index))
        name = str(item["name"])
        style_where = "%s: стиль «%s»" % (where, name)
        style_type = str(item.get("type") or "paragraph")
        if style_type not in STYLE_TYPES:
            raise StandardError("%s: type должен быть %s" % (style_where, ", ".join(STYLE_TYPES)))
        basis = parse_basis(str(item.get("basis", "proposal")), standard_basis, style_where)
        props_data = {key: value for key, value in item.items() if key not in STYLE_META}
        word_name, style_id = BUILTIN_STYLES.get(name, (name, style_id_for(name)))
        if style_id in seen_ids:
            raise StandardError("%s: стиль с таким именем уже есть («%s»)" % (style_where, seen_ids[style_id]))
        seen_ids[style_id] = name
        standard.styles.append(Style(
            name=name, type=style_type, role=str(item.get("role") or ""),
            based_on=str(item["based_on"]) if item.get("based_on") else None,
            next=str(item["next"]) if item.get("next") else None,
            note=str(item.get("note") or ""),
            props=read_values(props_data, STYLE_KEYS, basis, standard_basis, style_where),
            rules=read_rules(item.get("rules"), standard_basis, style_where),
            word_name=word_name, style_id=style_id, builtin=name in BUILTIN_STYLES,
        ))
    check_references(standard)
    return standard


def check_references(standard: Standard) -> None:
    names = standard.by_name
    if "Обычный" not in names:
        raise StandardError("%s: нет стиля «Обычный»" % standard.path.name)
    for style in standard.styles:
        where = "%s: стиль «%s»" % (standard.path.name, style.name)
        for label, target in (("based_on", style.based_on), ("next", style.next)):
            if target and target not in names:
                raise StandardError("%s: %s ссылается на стиль «%s», которого нет в файле" % (where, label, target))
        if "align" in style.props and style.props["align"].value not in ALIGNMENTS:
            raise StandardError("%s: align должен быть %s" % (where, ", ".join(ALIGNMENTS)))
        if "numbering" in style.props:
            list_id, level = parse_numbering(str(style.props["numbering"].value), where)
            if list_id not in standard.numbering or not 1 <= level <= len(standard.numbering[list_id]):
                raise StandardError("%s: нумерации %r нет в разделе numbering" % (where, style.props["numbering"].value))
        if "tabs" in style.props:
            parse_tabs(str(style.props["tabs"].value), where)
        if "borders" in style.props:
            parse_borders(str(style.props["borders"].value), where)
    chain_check = {}
    for style in standard.styles:
        seen = set()
        current = style
        while current is not None and current.based_on:
            if current.name in seen:
                raise StandardError("%s: стиль «%s» наследует сам себя по кругу" % (standard.path.name, style.name))
            seen.add(current.name)
            current = names.get(current.based_on)
        chain_check[style.name] = seen


def parse_numbering(text: str, where: str = "") -> tuple[str, int]:
    parts = text.split()
    if len(parts) != 2 or not parts[1].isdigit():
        raise StandardError("%s: numbering пишется как '<список> <уровень>'" % where)
    return parts[0], int(parts[1])


def parse_tabs(text: str, where: str = "") -> list[tuple[str, float, bool]]:
    tabs = []
    for chunk in [part.strip() for part in text.split(";") if part.strip()]:
        parts = chunk.split()
        if len(parts) not in (2, 3) or parts[0] not in TAB_KINDS or (len(parts) == 3 and parts[2] != "dot"):
            raise StandardError("%s: табуляция пишется как 'right 16.5 dot; center 8.25'" % where)
        tabs.append((parts[0], float(parts[1].replace(",", ".")), len(parts) == 3))
    return tabs


def parse_borders(text: str, where: str = "") -> float:
    parts = text.split()
    if len(parts) != 2 or parts[0] != "single":
        raise StandardError("%s: линии пишутся как 'single 0.5'" % where)
    return float(parts[1].replace(",", "."))


def effective_props(standard: Standard, style: Style) -> dict[str, Value]:
    chain = []
    current = style
    while current is not None:
        chain.append(current)
        current = standard.by_name.get(current.based_on) if current.based_on else None
    props: dict[str, Value] = {}
    if style.type == "paragraph":
        props.update(standard.defaults)
    for item in reversed(chain):
        props.update(item.props)
    return props


# ----------------------------------------------------------------- XML parts

def el(tag: str, attrs: dict | None = None, children=()) -> etree._Element:
    node = etree.Element(W + tag)
    for key, value in (attrs or {}).items():
        node.set(W + key, str(value))
    for child in children:
        if child is not None:
            node.append(child)
    return node


def flag(tag: str, on: bool) -> etree._Element:
    return el(tag) if on else el(tag, {"val": "0"})


def cm(value: float) -> int:
    return int(round(value * TWIPS_PER_CM))


def points(value: float) -> int:
    return int(round(value * 20))


def run_properties(props: dict[str, Value]) -> etree._Element | None:
    children = []
    if "font" in props:
        font = props["font"].value
        children.append(el("rFonts", {"ascii": font, "hAnsi": font, "cs": font, "eastAsia": font}))
    if "bold" in props:
        children += [flag("b", props["bold"].value), flag("bCs", props["bold"].value)]
    if "italic" in props:
        children += [flag("i", props["italic"].value), flag("iCs", props["italic"].value)]
    if "caps" in props:
        children.append(flag("caps", props["caps"].value))
    if "color" in props:
        color = str(props["color"].value)
        children.append(el("color", {"val": "auto" if color.lower() == "auto" else color.upper()}))
    if "size" in props:
        half_points = int(round(props["size"].value * 2))
        children += [el("sz", {"val": half_points}), el("szCs", {"val": half_points})]
    if "underline" in props:
        children.append(el("u", {"val": "single" if props["underline"].value else "none"}))
    if "vert_align" in props:
        children.append(el("vertAlign", {"val": props["vert_align"].value}))
    return el("rPr", children=children) if children else None


def paragraph_properties(standard: Standard, style: Style) -> etree._Element | None:
    props = style.props
    children = []
    for key, tag in (("keep_next", "keepNext"), ("keep_lines", "keepLines"), ("page_break_before", "pageBreakBefore")):
        if key in props:
            children.append(flag(tag, props[key].value))
    if "numbering" in props:
        list_id, level = parse_numbering(str(props["numbering"].value))
        children.append(el("numPr", children=[el("ilvl", {"val": level - 1}),
                                              el("numId", {"val": standard.numbering_ids[list_id]})]))
    if "tabs" in props:
        tabs = []
        for kind, position, dotted in parse_tabs(str(props["tabs"].value)):
            attrs = {"val": kind, "pos": cm(position)}
            if dotted:
                attrs["leader"] = "dot"
            tabs.append(el("tab", attrs))
        children.append(el("tabs", children=tabs))
    if "no_hyphenation" in props:
        children.append(flag("suppressAutoHyphens", props["no_hyphenation"].value))
    spacing = {}
    if "before" in props:
        spacing["before"] = points(props["before"].value)
    if "after" in props:
        spacing["after"] = points(props["after"].value)
    if "line" in props:
        spacing["line"] = int(round(240 * props["line"].value))
        spacing["lineRule"] = "auto"
    if spacing:
        children.append(el("spacing", spacing))
    indent = {}
    if "left" in props:
        indent["left"] = cm(props["left"].value)
    if "first_line" in props:
        first = props["first_line"].value
        indent["hanging" if first < 0 else "firstLine"] = cm(abs(first))
    if indent:
        children.append(el("ind", indent))
    if "align" in props:
        children.append(el("jc", {"val": ALIGNMENTS[props["align"].value]}))
    if "outline" in props:
        children.append(el("outlineLvl", {"val": props["outline"].value - 1}))
    return el("pPr", children=children) if children else None


def table_properties(style: Style) -> etree._Element:
    children = [el("tblInd", {"w": 0, "type": "dxa"})]
    if "borders" in style.props:
        size = int(round(parse_borders(str(style.props["borders"].value)) * 8))
        children.append(el("tblBorders", children=[
            el(side, {"val": "single", "sz": size, "space": 0, "color": "auto"})
            for side in ("top", "left", "bottom", "right", "insideH", "insideV")
        ]))
    children.append(el("tblCellMar", children=[
        el("top", {"w": 0, "type": "dxa"}), el("left", {"w": 108, "type": "dxa"}),
        el("bottom", {"w": 0, "type": "dxa"}), el("right", {"w": 108, "type": "dxa"}),
    ]))
    return el("tblPr", children=children)


def word_default_styles() -> list[etree._Element]:
    hidden = lambda: [el("semiHidden"), el("unhideWhenUsed")]  # noqa: E731
    return [
        el("style", {"type": "character", "default": "1", "styleId": "DefaultParagraphFont"},
           [el("name", {"val": "Default Paragraph Font"}), el("uiPriority", {"val": 1}), *hidden()]),
        el("style", {"type": "table", "default": "1", "styleId": "TableNormal"},
           [el("name", {"val": "Normal Table"}), el("uiPriority", {"val": 99}), *hidden(),
            el("tblPr", children=[el("tblInd", {"w": 0, "type": "dxa"}), el("tblCellMar", children=[
                el("top", {"w": 0, "type": "dxa"}), el("left", {"w": 108, "type": "dxa"}),
                el("bottom", {"w": 0, "type": "dxa"}), el("right", {"w": 108, "type": "dxa"})])])]),
        el("style", {"type": "numbering", "default": "1", "styleId": "NoList"},
           [el("name", {"val": "No List"}), el("uiPriority", {"val": 99}), *hidden()]),
    ]


def styles_xml(standard: Standard) -> etree._Element:
    root = etree.Element(W + "styles", nsmap=NSMAP)
    defaults = standard.defaults
    font = defaults["font"].value if "font" in defaults else "Times New Roman"
    half_points = int(round((defaults["size"].value if "size" in defaults else 12) * 2))
    line = int(round(240 * (defaults["line"].value if "line" in defaults else 1)))
    root.append(el("docDefaults", children=[
        el("rPrDefault", children=[el("rPr", children=[
            el("rFonts", {"ascii": font, "hAnsi": font, "cs": font, "eastAsia": font}),
            el("sz", {"val": half_points}), el("szCs", {"val": half_points}),
            el("lang", {"val": "ru-RU", "eastAsia": "ru-RU", "bidi": "ar-SA"}),
        ])]),
        el("pPrDefault", children=[el("pPr", children=[el("spacing", {"after": 0, "line": line, "lineRule": "auto"})])]),
    ]))
    # Built-in styles outside the set stay out of the gallery and the style pane.
    root.append(el("latentStyles", {"defLockedState": 0, "defUIPriority": 99, "defSemiHidden": 1,
                                    "defUnhideWhenUsed": 1, "defQFormat": 0, "count": 376}))
    by_name = standard.by_name
    for priority, style in enumerate(standard.styles, 1):
        attrs = {"type": style.type}
        if style.name == "Обычный":
            attrs["default"] = "1"
        if not style.builtin:
            attrs["customStyle"] = "1"
        attrs["styleId"] = style.style_id
        node = el("style", attrs, [el("name", {"val": style.word_name})])
        if style.based_on:
            node.append(el("basedOn", {"val": by_name[style.based_on].style_id}))
        elif style.type == "character":
            node.append(el("basedOn", {"val": "DefaultParagraphFont"}))
        elif style.type == "table":
            node.append(el("basedOn", {"val": "TableNormal"}))
        if style.next:
            node.append(el("next", {"val": by_name[style.next].style_id}))
        node.append(el("uiPriority", {"val": priority}))
        if style.type == "paragraph" and style.name not in OUTSIDE_GALLERY:
            node.append(el("qFormat"))
        if style.type == "paragraph":
            parts = [paragraph_properties(standard, style), run_properties(style.props)]
        elif style.type == "character":
            parts = [run_properties(style.props)]
        else:
            parts = [el("pPr", children=[el("spacing", {"after": 0, "line": 240, "lineRule": "auto"})]),
                     table_properties(style)]
        for part in parts:
            if part is not None:
                node.append(part)
        root.append(node)
    for node in word_default_styles():
        root.append(node)
    return root


def numbering_xml(standard: Standard) -> etree._Element:
    root = etree.Element(W + "numbering", nsmap=NSMAP)
    owners: dict[tuple[str, int], Style] = {}
    for style in standard.styles:
        if "numbering" in style.props:
            owners.setdefault(parse_numbering(str(style.props["numbering"].value)), style)
    nums = []
    for index, (list_id, levels) in enumerate(standard.numbering.items()):
        multilevel = len(levels) > 1
        abstract = el("abstractNum", {"abstractNumId": index},
                      [el("multiLevelType", {"val": "multilevel" if multilevel else "singleLevel"})])
        first_owner = owners.get((list_id, 1))
        for level in range(9 if multilevel else len(levels)):
            spec = levels[level] if level < len(levels) else {
                "format": "decimal", "text": ".".join("%%%d" % (k + 1) for k in range(level + 1))}
            owner = owners.get((list_id, level + 1)) or first_owner
            props = effective_props(standard, owner) if owner else dict(standard.defaults)
            first = props["first_line"].value if "first_line" in props else 0
            left = props["left"].value if "left" in props else 0
            lvl = el("lvl", {"ilvl": level}, [el("start", {"val": 1}), el("numFmt", {"val": spec["format"]})])
            if owner is not None and owners.get((list_id, level + 1)) is owner:
                lvl.append(el("pStyle", {"val": owner.style_id}))
            lvl.append(el("suff", {"val": "space"}))
            lvl.append(el("lvlText", {"val": spec["text"]}))
            lvl.append(el("lvlJc", {"val": "left"}))
            indent = {"left": cm(left)}
            indent["hanging" if first < 0 else "firstLine"] = cm(abs(first))
            lvl.append(el("pPr", children=[el("ind", indent)]))
            if spec["format"] == "bullet" and "font" in props:
                font = props["font"].value
                lvl.append(el("rPr", children=[el("rFonts", {"ascii": font, "hAnsi": font, "cs": font, "hint": "default"})]))
            abstract.append(lvl)
        root.append(abstract)
        nums.append(el("num", {"numId": index + 1}, [el("abstractNumId", {"val": index})]))
    for num in nums:
        root.append(num)
    return root


def settings_xml() -> etree._Element:
    uri = "http://schemas.microsoft.com/office/word"
    compat = [("compatibilityMode", "15"), ("overrideTableStyleFontSizeAndJustification", "1"),
              ("enableOpenTypeFeatures", "1"), ("doNotFlipMirrorIndents", "1"),
              ("differentiateMultirowTableHeaders", "1")]
    root = etree.Element(W + "settings", nsmap=NSMAP)
    for child in (
        el("zoom", {"percent": 100}),
        el("defaultTabStop", {"val": 709}),
        el("characterSpacingControl", {"val": "doNotCompress"}),
        el("compat", children=[el("compatSetting", {"name": name, "uri": uri, "val": value}) for name, value in compat]),
        el("themeFontLang", {"val": "ru-RU"}),
        el("decimalSymbol", {"val": ","}),
        el("listSeparator", {"val": ";"}),
    ):
        root.append(child)
    return root


def section_properties(standard: Standard) -> etree._Element:
    page = standard.page
    top, bottom, left, right = (int(round(value * TWIPS_PER_CM / 10)) for value in page["margins_mm"].value)
    header = int(round((page["header_mm"].value if "header_mm" in page else 12.5) * TWIPS_PER_CM / 10))
    footer = int(round((page["footer_mm"].value if "footer_mm" in page else 12.5) * TWIPS_PER_CM / 10))
    reference = el("footerReference", {"type": "default"})
    reference.set("{%s}id" % R_NS, "rIdFooter1")
    section = el("sectPr", children=[
        reference,
        el("pgSz", {"w": A4_TWIPS[0], "h": A4_TWIPS[1]}),
        el("pgMar", {"top": top, "right": right, "bottom": bottom, "left": left,
                     "header": header, "footer": footer, "gutter": 0}),
        el("cols", {"space": 708}),
    ])
    if "first_page_number" in page and not page["first_page_number"].value:
        section.append(el("titlePg"))
    section.append(el("docGrid", {"linePitch": 360}))
    return section


def text_run(text: str, style: Style | None = None) -> list[etree._Element]:
    runs = []
    for index, line in enumerate(text.split("\n")):
        run = el("r")
        if style is not None:
            run.append(el("rPr", children=[el("rStyle", {"val": style.style_id})]))
        if index:
            run.append(el("br"))
        for piece_index, piece in enumerate(line.split("\t")):
            if piece_index:
                run.append(el("tab"))
            if piece:
                node = el("t")
                node.text = piece
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                run.append(node)
        runs.append(run)
    return runs


def field_run(kind: str, style: Style | None = None, instruction: str = "") -> etree._Element:
    run = el("r")
    if style is not None:
        run.append(el("rPr", children=[el("rStyle", {"val": style.style_id})]))
    if kind == "instr":
        node = el("instrText")
        node.text = instruction
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        run.append(node)
    else:
        run.append(el("fldChar", {"fldCharType": kind}))
    return run


def footer_xml(standard: Standard) -> etree._Element:
    root = etree.Element(W + "ftr", nsmap=NSMAP)
    footer_style = standard.by_name.get("Нижний колонтитул")
    number_style = standard.by_name.get("Номер страницы")
    alignment = str(standard.page["page_number"].value) if "page_number" in standard.page else "center"
    ppr = el("pPr")
    if footer_style is not None:
        ppr.append(el("pStyle", {"val": footer_style.style_id}))
    ppr.append(el("jc", {"val": ALIGNMENTS.get(alignment, "center")}))
    paragraph = el("p", children=[ppr])
    paragraph.append(field_run("begin", number_style))
    paragraph.append(field_run("instr", number_style, " PAGE "))
    paragraph.append(field_run("separate", number_style))
    paragraph.extend(text_run("1", number_style))
    paragraph.append(field_run("end", number_style))
    root.append(paragraph)
    return root


def xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def package_parts(standard: Standard, body_blocks: list[etree._Element], template: bool) -> dict[str, bytes]:
    document = etree.Element(W + "document", nsmap=NSMAP)
    body = el("body", children=body_blocks)
    body.append(section_properties(standard))
    document.append(body)
    main_type = TEMPLATE_CONTENT_TYPE if template else DOCUMENT_CONTENT_TYPE
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="%s"/>'
        '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
        '<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>'
        '<Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>'
        '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
        '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
        '</Types>' % main_type
    )
    package_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="%sofficeDocument" Target="word/document.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        '<Relationship Id="rId3" Type="%sextended-properties" Target="docProps/app.xml"/>'
        '</Relationships>' % (rel, rel)
    )
    document_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdStyles" Type="%sstyles" Target="styles.xml"/>'
        '<Relationship Id="rIdNumbering" Type="%snumbering" Target="numbering.xml"/>'
        '<Relationship Id="rIdSettings" Type="%ssettings" Target="settings.xml"/>'
        '<Relationship Id="rIdFooter1" Type="%sfooter" Target="footer1.xml"/>'
        '</Relationships>' % (rel, rel, rel, rel)
    )
    core = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        '<dc:title>%s</dc:title></cp:coreProperties>' % escape(standard.title)
    )
    app = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
        '<Application>Audion DocFlow</Application></Properties>'
    )
    return {
        "[Content_Types].xml": content_types.encode("utf-8"),
        "_rels/.rels": package_rels.encode("utf-8"),
        "docProps/core.xml": core.encode("utf-8"),
        "docProps/app.xml": app.encode("utf-8"),
        "word/document.xml": xml_bytes(document),
        "word/_rels/document.xml.rels": document_rels.encode("utf-8"),
        "word/styles.xml": xml_bytes(styles_xml(standard)),
        "word/numbering.xml": xml_bytes(numbering_xml(standard)),
        "word/settings.xml": xml_bytes(settings_xml()),
        "word/footer1.xml": xml_bytes(footer_xml(standard)),
    }


def write_package(path: Path, parts: dict[str, bytes]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


# ----------------------------------------------------------------- sample

class SampleWriter:
    """Body blocks of a sample document; every paragraph names a style of the standard."""

    def __init__(self, standard: Standard) -> None:
        self.standard = standard
        self.blocks: list[etree._Element] = []

    def style(self, name: str) -> Style:
        try:
            return self.standard.by_name[name]
        except KeyError:
            raise StandardError("%s: образцу нужен стиль «%s»" % (self.standard.path.name, name)) from None

    def paragraph(self, style_name: str, text: str = "", runs: list[tuple[str, str | None]] | None = None) -> etree._Element:
        node = el("p", children=[el("pPr", children=[el("pStyle", {"val": self.style(style_name).style_id})])])
        for piece, run_style in runs or [(text, None)]:
            node.extend(text_run(piece, self.style(run_style) if run_style else None))
        return node

    def add(self, style_name: str, text: str = "", runs: list[tuple[str, str | None]] | None = None) -> None:
        self.blocks.append(self.paragraph(style_name, text, runs))

    def items(self, style_name: str, texts: list[str]) -> None:
        for text in texts:
            self.add(style_name, text)

    def toc(self, entries: list[tuple[str, str, str]]) -> None:
        instruction = " TOC %so \"1-3\" %sh %sz %su " % (BACKSLASH, BACKSLASH, BACKSLASH, BACKSLASH)
        for index, (style_name, text, page) in enumerate(entries):
            node = self.paragraph(style_name)
            if index == 0:
                node.append(field_run("begin"))
                node.append(field_run("instr", instruction=instruction))
                node.append(field_run("separate"))
            node.extend(text_run("%s\t%s" % (text, page)))
            if index == len(entries) - 1:
                node.append(field_run("end"))
            self.blocks.append(node)

    def table(self, header: list[str], rows: list[list[str]], widths: list[float], aligns: list[str]) -> None:
        width_twips = cm(self.standard.text_width_cm())
        columns = [int(width_twips * share / sum(widths)) for share in widths]
        grid = self.style("Сетка таблицы")
        table = el("tbl", children=[
            el("tblPr", children=[el("tblStyle", {"val": grid.style_id}), el("tblW", {"w": sum(columns), "type": "dxa"}),
                                  el("tblLayout", {"type": "fixed"})]),
            el("tblGrid", children=[el("gridCol", {"w": column}) for column in columns]),
        ])
        header_row = el("tr", children=[el("trPr", children=[el("tblHeader")])])
        for text, column in zip(header, columns):
            header_row.append(self.cell(text, column, "Табличный_заголовки"))
        table.append(header_row)
        for row in rows:
            node = el("tr")
            for text, column, align in zip(row, columns, aligns):
                node.append(self.cell(text, column, "Табличный_центр" if align == "center" else "Табличный_слева"))
            table.append(node)
        self.blocks.append(table)
        self.add("Абзац")  # the empty line after a table

    def cell(self, text: str, width: int, style_name: str) -> etree._Element:
        return el("tc", children=[el("tcPr", children=[el("tcW", {"w": width, "type": "dxa"})]),
                                  self.paragraph(style_name, text)])


def notes_sample(writer: SampleWriter) -> None:
    add = writer.add
    add("Заголовок оглавления", "Содержание")
    writer.toc([
        ("Оглавление 1", "1 Общие положения", "2"),
        ("Оглавление 2", "1.1 Назначение документа", "2"),
        ("Оглавление 2", "1.2 Исходные данные", "2"),
        ("Оглавление 1", "2 Основные решения", "3"),
        ("Оглавление 2", "2.1 Показатели", "3"),
        ("Оглавление 3", "2.1.1 Сводные показатели", "3"),
        ("Оглавление 1", "Приложение 1 Перечень исходных материалов", "4"),
    ])
    add("Заголовок 1", "Общие положения")
    add("Заголовок 2", "Назначение документа")
    add("Абзац", "Образец показывает стили оформления записки. Текст условный: важны шрифт, отступы, "
                 "интервалы и нумерация, а не содержание.")
    add("Абзац", "Основной текст набирается стилем «Абзац». Отступ первой строки, выравнивание по ширине "
                 "и интервал после абзаца заданы в стиле, вручную их ставить не нужно.")
    add("Абзац", "Записка включает:")
    writer.items("Маркированный список", ["исходные данные;", "расчётные показатели;", "выводы и предложения."])
    add("Абзац", "Работа выполняется в три этапа:")
    writer.items("Нумерованный список", ["сбор исходных данных;", "расчёт показателей;", "подготовка выводов."])
    add("Абзац", "Если на элемент перечисления есть ссылка в тексте, вместо цифр ставят буквы:")
    writer.items("Нумерованный список 2", ["первый вариант размещения;", "второй вариант размещения."])
    add("Заголовок 2", "Исходные данные")
    add("Абзац", "Основные показатели приведены в таблице 1.")
    add("Название таблицы", "Таблица 1 – Основные показатели")
    writer.table(["№ п/п", "Показатель", "Единица измерения", "Значение"],
                 [["1", "Площадь территории", "га", "125,4"],
                  ["2", "Численность населения", "тыс. чел.", "12,8"],
                  ["3", "Протяжённость улично-дорожной сети", "км", "46,2"]],
                 [1, 5, 2, 2], ["center", "left", "center", "center"])
    add("Заголовок 1", "Основные решения")
    add("Заголовок 2", "Показатели")
    add("Заголовок 3", "Сводные показатели")
    add("Абзац", "Сводные показатели собираются по разделам записки и сверяются с исходными данными.")
    add("Заголовок 4", "Уточнение показателей")
    add("Абзац", "Схема размещения объектов показана на рисунке 1.")
    add("Рисунок", "[ место для рисунка ]")
    add("Название объекта", "Рисунок 1 – Схема размещения объектов")
    add("Абзац", runs=[("Порядок уточнения показателей описан в ", None), ("разделе 1", "Гиперссылка"), (".", None)])
    add("Заголовок приложения", "ПРИЛОЖЕНИЕ 1\nПеречень исходных материалов")
    add("Абзац", "Перечень составляется по мере поступления материалов.")


def nir_sample(writer: SampleWriter) -> None:
    add = writer.add
    add("Заголовок структурного элемента", "Реферат")
    add("Абзац", "Отчёт 12 с., 1 рис., 1 табл., 2 источн., 1 прил.")
    add("Абзац без отступа", "ИСХОДНЫЕ ДАННЫЕ, ПОКАЗАТЕЛИ ТЕРРИТОРИИ, МЕТОДИКА РАСЧЁТА, ПРЕДЛОЖЕНИЯ")
    add("Абзац", "Образец показывает стили оформления отчёта о НИР. Текст условный: важны шрифт, отступы, "
                 "интервалы и нумерация, а не содержание.")
    add("Заголовок структурного элемента", "Содержание")
    writer.toc([
        ("Оглавление 1", "Термины и определения", "3"),
        ("Оглавление 1", "Введение", "4"),
        ("Оглавление 1", "1 Анализ исходных данных", "5"),
        ("Оглавление 2", "1.1 Состав исходных данных", "5"),
        ("Оглавление 3", "1.1.1 Показатели территории", "6"),
        ("Оглавление 1", "Заключение", "8"),
        ("Оглавление 1", "Список использованных источников", "9"),
        ("Оглавление 1", "Приложение А (справочное) Перечень исходных материалов", "10"),
    ])
    add("Заголовок структурного элемента", "Термины и определения")
    add("Абзац без отступа", "Показатель – величина, которая характеризует состояние территории")
    add("Абзац без отступа", "Территория – часть земной поверхности в установленных границах")
    add("Заголовок структурного элемента", "Введение")
    add("Абзац", "Основной текст набирается стилем «Абзац»: Times New Roman 12 пт, полуторный интервал, "
                 "абзацный отступ 1,25 см.")
    add("Заголовок 1", "Анализ исходных данных")
    add("Заголовок 2", "Состав исходных данных")
    add("Абзац", "Исходные данные включают [1]:")
    writer.items("Маркированный список", ["картографические материалы,", "статистические данные,", "нормативные документы."])
    add("Абзац", "Работа выполнялась в два этапа:")
    writer.items("Нумерованный список", ["сбор и проверка данных;", "расчёт показателей."])
    add("Абзац", "Для расчёта рассмотрены варианты:")
    writer.items("Нумерованный список 2", ["расчёт по фактическим данным;", "расчёт по нормативам."])
    add("Заголовок 3", "Показатели территории")
    add("Абзац", "Площадь участка рассчитывают по формуле (1)")
    add("Формула", "\tS = a × b\t(1)")
    add("Абзац без отступа", "где S – площадь участка, м²;")
    add("Абзац без отступа", "a – длина участка, м;")
    add("Абзац без отступа", "b – ширина участка, м.")
    add("Абзац", "Результаты расчёта приведены в таблице 1.")
    add("Название таблицы", "Таблица 1 – Показатели территории")
    writer.table(["Показатель", "Единица измерения", "Значение"],
                 [["Площадь территории", "га", "125,4"], ["Численность населения", "тыс. чел.", "12,8"]],
                 [5, 2, 2], ["left", "center", "center"])
    add("Абзац", "Схема территории показана на рисунке 1.")
    add("Рисунок", "[ место для рисунка ]")
    add("Название объекта", "Рисунок 1 – Схема территории")
    add("Заголовок структурного элемента", "Заключение")
    add("Абзац", "В заключении кратко излагают выводы по результатам работы.")
    add("Заголовок структурного элемента", "Список использованных источников")
    writer.items("Список литературы", [
        "ГОСТ 7.32-2017 Система стандартов по информации, библиотечному и издательскому делу. "
        "Отчёт о научно-исследовательской работе. Структура и правила оформления.",
        "ГОСТ 8.417-2024 Государственная система обеспечения единства измерений. Единицы величин.",
    ])
    add("Заголовок приложения", "ПРИЛОЖЕНИЕ А\n(справочное)\nПеречень исходных материалов")
    add("Абзац", "Перечень составляется по мере поступления материалов.")


SAMPLES = {"notes": notes_sample, "nir_gost_7_32": nir_sample}


def build_template(standard: Standard, path: Path) -> Path:
    blocks = [SampleWriter(standard).paragraph("Абзац")]
    write_package(path, package_parts(standard, blocks, template=True))
    return path


def build_sample(standard: Standard, path: Path) -> Path:
    fill = SAMPLES.get(standard.id)
    if fill is None:
        raise StandardError("%s: для стандарта %r нет образца" % (standard.path.name, standard.id))
    writer = SampleWriter(standard)
    fill(writer)
    write_package(path, package_parts(standard, writer.blocks, template=False))
    return path


# ----------------------------------------------------------------- description

LABELS = [
    ("font", "Шрифт"), ("size", "Кегль"), ("bold", "Полужирный"), ("italic", "Курсив"),
    ("caps", "Все прописные"), ("underline", "Подчёркивание"), ("color", "Цвет текста"),
    ("vert_align", "Положение знака"), ("align", "Выравнивание"), ("first_line", "Отступ первой строки"),
    ("left", "Отступ слева"), ("line", "Межстрочный интервал"), ("before", "Интервал перед"),
    ("after", "Интервал после"), ("tabs", "Табуляция"), ("numbering", "Нумерация"),
    ("keep_next", "Не отрывать от следующего"), ("keep_lines", "Не разрывать абзац"),
    ("page_break_before", "С новой страницы"), ("no_hyphenation", "Переносы слов"),
    ("outline", "Уровень заголовка"), ("borders", "Линии"),
]
PAGE_LABELS = [("size", "Лист"), ("margins_mm", "Поля"), ("header_mm", "Верхний колонтитул"),
               ("footer_mm", "Нижний колонтитул"), ("page_number", "Номер страницы"),
               ("first_page_number", "Номер на первом листе")]
ALIGN_TEXT = {"left": "по левому краю", "center": "по центру", "right": "по правому краю", "both": "по ширине"}
TYPE_TEXT = {"paragraph": "абзац", "character": "знак", "table": "таблица"}


def ru_number(value: float) -> str:
    text = ("%.2f" % value).rstrip("0").rstrip(".")
    return text.replace(".", ",")


def basis_text(standard: Standard, basis: str) -> str:
    if basis == "proposal":
        return "Предложение"
    if basis == "standard":
        return standard.standard_basis
    clauses = [part.strip() for part in basis[len("gost "):].split(",") if part.strip()]
    return "%s, %s %s" % (standard.gost_name, "пп." if len(clauses) > 1 else "п.", ", ".join(clauses))


def numbering_text(standard: Standard, value: str) -> str:
    list_id, level = parse_numbering(value)
    spec = standard.numbering[list_id][level - 1]
    if spec["format"] == "bullet":
        return "маркер «%s»" % spec["text"]
    sample = re.sub(r"%\d", lambda _: NUMBER_FORMATS[spec["format"]], spec["text"])
    return "автоматическая: «%s»" % sample


def value_text(standard: Standard, key: str, value: object) -> str:
    if key in ("bold", "italic", "caps", "underline", "keep_next", "keep_lines", "page_break_before"):
        return "да" if value else "нет"
    if key == "no_hyphenation":
        return "запрещены" if value else "разрешены"
    if key == "size":
        return "%s пт" % ru_number(value)
    if key in ("before", "after"):
        return "%s пт" % ru_number(value)
    if key in ("first_line", "left"):
        if not value:
            return "нет"
        return ("выступ %s см" if value < 0 else "%s см") % ru_number(abs(value))
    if key == "line":
        return {1.0: "одинарный", 1.5: "полуторный", 2.0: "двойной"}.get(float(value), "множитель " + ru_number(value))
    if key == "align":
        return ALIGN_TEXT[value]
    if key == "color":
        text = str(value).lower()
        return {"auto": "авто (чёрный)", "000000": "чёрный", "0000ff": "синий"}.get(text, "#" + text.upper())
    if key == "vert_align":
        return {"superscript": "надстрочный", "subscript": "подстрочный"}.get(value, str(value))
    if key == "tabs":
        return "; ".join("%s %s см%s" % (TAB_KINDS[kind], ru_number(position), " с отточием" if dotted else "")
                         for kind, position, dotted in parse_tabs(str(value)))
    if key == "numbering":
        return numbering_text(standard, str(value))
    if key == "borders":
        return "одинарные %s пт со всех сторон и внутри" % ru_number(parse_borders(str(value)))
    return str(value)


def page_value_text(key: str, value: object) -> str:
    if key == "size":
        return "%s книжный" % value
    if key == "margins_mm":
        top, bottom, left, right = value
        return "верхнее %s мм, нижнее %s мм, левое %s мм, правое %s мм" % tuple(ru_number(v) for v in (top, bottom, left, right))
    if key in ("header_mm", "footer_mm"):
        return "%s мм от края листа" % ru_number(value)
    if key == "page_number":
        return {"right": "внизу справа", "center": "внизу по центру", "left": "внизу слева"}.get(value, str(value))
    if key == "first_page_number":
        return "ставится" if value else "не ставится"
    return str(value)


def build_description(standards: list[Standard], path: Path, date_text: str) -> tuple[int, int]:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    widths = (Cm(3.6), Cm(6.3), Cm(4.0), Cm(3.6))
    font_name = "Tahoma"
    rows_total = proposals_total = 0

    def set_font(owner, size: float, bold: bool | None = None, color: RGBColor | None = None) -> None:
        owner.font.name = font_name
        owner.font.size = Pt(size)
        if bold is not None:
            owner.font.bold = bold
        if color is not None:
            owner.font.color.rgb = color
        rpr = owner.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for key in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(key)) is not None:
                del fonts.attrib[qn(key)]
        for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(key), font_name)

    def spaced(paragraph, before: float = 0, after: float = 0):
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        return paragraph

    def shade(cell, fill: str) -> None:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shd)

    def cell_text(cell, text: str, bold: bool = False) -> None:
        run = spaced(cell.paragraphs[0]).add_run(text)
        run.font.size = Pt(10)
        run.font.bold = bold

    def keep_row(row) -> None:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

    def new_table(document, titles: tuple[str, ...], column_widths):
        table = document.add_table(rows=1, cols=len(titles))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        for grid_col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), column_widths):
            grid_col.set(qn("w:w"), str(int(width.twips)))
        for cell, title, width in zip(table.rows[0].cells, titles, column_widths):
            cell.width = width
            cell_text(cell, title, bold=True)
            shade(cell, "D9D9D9")
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        keep_row(table.rows[0])
        return table

    def add_rows(table, rows: list[tuple[str, str, str]], column_widths) -> None:
        nonlocal rows_total, proposals_total
        for param, value, basis in rows:
            row = table.add_row()
            keep_row(row)
            cells = row.cells
            for index, (cell, text, width) in enumerate(zip(cells, (param, value, basis, ""), column_widths)):
                cell.width = width
                cell_text(cell, text)
                if index == 2 and text == "Предложение":
                    shade(cell, "FFF2CC")
            rows_total += 1
            proposals_total += basis == "Предложение"

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = Cm(2)
    section.right_margin = Cm(1.5)
    normal = document.styles["Normal"]
    set_font(normal, 11)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    normal.element.get_or_add_rPr().append(lang)
    for level, size, before in ((1, 15, 12), (2, 12, 10)):
        heading = document.styles["Heading %d" % level]
        set_font(heading, size, bold=True, color=RGBColor(0, 0, 0))
        heading.paragraph_format.space_before = Pt(before)
        heading.paragraph_format.space_after = Pt(4)
        heading.paragraph_format.keep_with_next = True

    title = spaced(document.add_paragraph(), 0, 4).add_run("Стили документов: записки и отчёт о НИР")
    title.font.size = Pt(16)
    title.font.bold = True
    spaced(document.add_paragraph("Черновик для согласования, %s" % date_text), 0, 10)
    spaced(document.add_paragraph(
        "Для каждого стиля указаны параметры, значение и основание. Правку пишите в колонку «Поправка»; "
        "пустая клетка означает, что значение подходит. Жёлтым отмечены предложения: их нужно подтвердить "
        "или заменить."), 0, 6)
    spaced(document.add_paragraph("Основания:"), 0, 2)
    legend = []
    for standard in standards:
        if standard.standard_basis:
            legend.append("«%s» — утверждённое значение;" % standard.standard_basis)
        if any(v.basis.startswith("gost") for s in standard.styles for v in s.props.values()):
            legend.append("«%s, п. …» — требование стандарта;" % standard.gost_name)
    legend.append("«Предложение» — значение предложено до уточнения.")
    for text in dict.fromkeys(legend):
        spaced(document.add_paragraph(text, style="List Bullet"), 0, 2)
    spaced(document.add_paragraph(
        "Состав и названия стилей — тоже предложение. Существующие документы источником значений не служили. "
        "Написание чисел, единиц, сокращений, адресов и инициалов — не стили, а правила текста; "
        "они описаны отдельно."), 6, 0)

    for number, standard in enumerate(standards, 1):
        document.add_heading("%d. %s" % (number, standard.title), level=1)
        spaced(document.add_paragraph(standard.lead), 0, 6)
        spaced(document.add_paragraph(
            "Стилей в шаблоне «%s»: %d — %d описаны ниже и 3 служебных стиля Word по умолчанию "
            "(%s)." % (standard.template, standard.style_count, len(standard.styles), ", ".join(WORD_DEFAULT_STYLES))), 0, 6)
        summary_widths = (Cm(5.5), Cm(9.5), Cm(2.5))
        summary = new_table(document, ("Стиль", "Для чего", "Тип"), summary_widths)
        for style in standard.styles:
            row = summary.add_row()
            keep_row(row)
            cells = row.cells
            for cell, text, width in zip(cells, (style.name, style.role, TYPE_TEXT[style.type]), summary_widths):
                cell.width = width
                cell_text(cell, text)

        document.add_heading("Страница", level=2)
        table = new_table(document, ("Параметр", "Значение", "Основание", "Поправка"), widths)
        page_rows = [(label, page_value_text(key, standard.page[key].value), basis_text(standard, standard.page[key].basis))
                     for key, label in PAGE_LABELS if key in standard.page]
        page_rows += [(param, value, basis_text(standard, basis)) for param, value, basis in standard.page_rules]
        add_rows(table, page_rows, widths)

        for style in standard.styles:
            document.add_heading("%s — %s" % (style.name, style.role) if style.role else style.name, level=2)
            if style.note:
                spaced(document.add_paragraph(style.note), 0, 4)
            if style.based_on:
                spaced(document.add_paragraph(
                    "Основан на стиле «%s»: параметры, которых нет в таблице, — как у него." % style.based_on), 0, 4)
            props = style.props
            rows = [(label, value_text(standard, key, props[key].value), basis_text(standard, props[key].basis))
                    for key, label in LABELS if key in props]
            if style.next:
                rows.append(("Следующий абзац", "«%s»" % style.next, "Предложение"))
            rows += [(param, value, basis_text(standard, basis)) for param, value, basis in style.rules]
            table = new_table(document, ("Параметр", "Значение", "Основание", "Поправка"), widths)
            add_rows(table, rows, widths)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return rows_total, proposals_total


# ----------------------------------------------------------------- CLI

def main() -> int:
    parser = argparse.ArgumentParser(description="Build Word templates, samples and the review description from formatting standards.")
    parser.add_argument("--standard", action="append", required=True, help="Standard YAML file; repeat for several")
    parser.add_argument("--outdir", default="docs/standards", help="Folder for templates and samples")
    parser.add_argument("--description", default="", help="Review description DOCX built from all the standards")
    parser.add_argument("--date", default="", help="Date printed in the description")
    args = parser.parse_args()

    try:
        standards = [load_standard(Path(path)) for path in args.standard]
    except StandardError as exc:
        print("[FAILED] %s" % exc)
        return 1
    outdir = Path(args.outdir)
    for standard in standards:
        low, high = CORPORATE_STYLE_RANGE
        marker = "OK" if low <= standard.style_count <= high else "WARN"
        template = build_template(standard, outdir / standard.template)
        sample = build_sample(standard, outdir / standard.sample)
        print("[%s] %s: styles %d (corporate set %d-%d)" % (marker, standard.path.name, standard.style_count, low, high))
        print("[OK] Template: %s" % template)
        print("[OK] Sample: %s" % sample)
    if args.description:
        import datetime
        date_text = args.date or datetime.date.today().strftime("%d.%m.%Y")
        rows, proposals = build_description(standards, Path(args.description), date_text)
        print("[OK] Description: %s (parameter rows %d, proposals %d)" % (args.description, rows, proposals))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
