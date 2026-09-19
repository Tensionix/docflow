"""Text hygiene lists every change it makes: the kind, before, after and the context."""
from __future__ import annotations

import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PYTHON = ROOT / "runtime" / "python.exe"
SCRIPT = ROOT / "system_core" / "docx_text_hygiene_fix.py"


class TextHygieneReportTests(unittest.TestCase):
    def test_report_lists_each_change_with_visible_spaces(self) -> None:
        tmp = Path(tempfile.mkdtemp(prefix="hygiene_report_"))
        self.addCleanup(shutil.rmtree, tmp, True)
        source = tmp / "input"
        source.mkdir()
        document = Document()
        document.add_paragraph("При сжигании газа и мазута : диоксид азота.")
        document.add_paragraph("Итого  12 котельных.")
        document.save(str(source / "note.docx"))
        report = tmp / "changes.md"
        result = subprocess.run(
            [str(PYTHON), str(SCRIPT), "--input", str(source), "--outdir", str(tmp / "out"),
             "--report", str(report), "--json-out", str(tmp / "changes.json")],
            capture_output=True, text=True, encoding="utf-8",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        text = report.read_text(encoding="utf-8")
        self.assertIn("## Правки по видам", text)
        self.assertIn("| Пробел перед пунктуацией | 1 |", text)
        self.assertIn("| Двойные пробелы | 1 |", text)
        self.assertIn("| Пробел перед пунктуацией | `\u00b7:` | `:` | При сжигании газа и мазута : диоксид азота. |", text)
        self.assertIn("| Двойные пробелы | `\u00b7\u00b7` | `\u00b7` |", text)
        payload = json.loads((tmp / "changes.json").read_text(encoding="utf-8"))
        classes = sorted(change["class_id"] for change in payload["files"][0]["findings"])
        self.assertEqual(classes, ["double_space", "space_before_punct"])
        self.assertTrue(report.with_suffix(".docx").is_file())


if __name__ == "__main__":
    unittest.main()
